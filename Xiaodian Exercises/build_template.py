"""Build the downloadable question-bank template for 小点刷题."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


OUT = Path("assets") / "question-bank-template.xlsx"

BRAND = "1F4A3B"
BRAND_DARK = "173A2E"
CREAM = "FFFDF8"
CREAM_2 = "F7F2E7"
LINE = "E2DBCE"
ACCENT = "DD6B43"
MUTED = "747E77"
GOLD = "F6EDD8"
SOFT_GREEN = "EAF4ED"
SOFT_RED = "FAEAE6"

thin = Side(style="thin", color=LINE)
cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = [
    "题型",
    "题干",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "正确答案",
    "解析",
    "分值",
    "分类",
    "难度",
]

rows = [
    [
        "单选",
        "“学而不思则罔，思而不学则殆”出自下列哪部典籍？",
        "《论语》",
        "《孟子》",
        "《大学》",
        "《中庸》",
        "",
        "A",
        "这句话出自《论语·为政》，强调学习与思考必须相互结合。只读书不思考会迷惑，只空想不读书也会陷入疑惑。",
        5,
        "传统文化",
        "简单",
    ],
    [
        "多选",
        "下列哪些属于 Python 的内置数据类型？",
        "list",
        "dict",
        "array",
        "tuple",
        "",
        "ABD",
        "list、dict、tuple 都是 Python 的内置类型。array 不是内置类型，需要通过 array 模块或 NumPy 等第三方库使用。",
        5,
        "计算机基础",
        "中等",
    ],
    [
        "判断",
        "在标准大气压下，纯水的沸点是 100°C。",
        "",
        "",
        "",
        "",
        "",
        "正确",
        "在标准大气压（101.325 kPa）下，纯水的沸点为 100°C。气压降低时，沸点也会随之降低。",
        5,
        "自然科学",
        "简单",
    ],
    [
        "单选",
        "在 HTTP 状态码中，404 通常表示什么？",
        "请求成功",
        "服务器内部错误",
        "请求的资源不存在",
        "页面被永久重定向",
        "",
        "C",
        "404 Not Found 表示服务器找不到客户端请求的资源。500 表示服务器内部错误，301/308 表示永久重定向。",
        5,
        "计算机基础",
        "简单",
    ],
]

wb = Workbook()
ws = wb.active
ws.title = "题库"
ws.sheet_properties.tabColor = BRAND
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A2"

# Column widths and row sizing
widths = [9, 42, 18, 18, 20, 20, 16, 12, 50, 8, 15, 9]
for index, width in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(index)].width = width

ws.row_dimensions[1].height = 32
ws.merge_cells("A1:L1")
title = ws["A1"]
title.value = "小点刷题 · 题库模板"
title.font = Font(name="微软雅黑", size=14, bold=True, color=CREAM)
title.fill = PatternFill("solid", fgColor=BRAND)
title.alignment = Alignment(horizontal="left", vertical="center", indent=1)

# Put actual headers on row 2; our importer searches the first 15 rows.
ws.row_dimensions[2].height = 28
for column, value in enumerate(headers, start=1):
    cell = ws.cell(row=2, column=column, value=value)
    cell.font = Font(name="微软雅黑", size=10, bold=True, color=CREAM)
    cell.fill = PatternFill("solid", fgColor=BRAND_DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = cell_border

for row_index, data in enumerate(rows, start=3):
    ws.row_dimensions[row_index].height = 70 if row_index == 3 else 58
    for column, value in enumerate(data, start=1):
        cell = ws.cell(row=row_index, column=column, value=value)
        cell.font = Font(name="微软雅黑", size=10, color="1F2D27")
        cell.border = cell_border
        cell.alignment = Alignment(
            horizontal="center" if column in (1, 8, 10, 11, 12) else "left",
            vertical="top" if column in (2, 9) else "center",
            wrap_text=column in (2, 3, 4, 5, 6, 7, 9),
        )
        if row_index % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=CREAM_2)

# Make the answer column easy to spot
for row_index in range(3, 3 + len(rows)):
    answer_cell = ws.cell(row=row_index, column=8)
    answer_cell.font = Font(name="微软雅黑", size=10, bold=True, color=ACCENT)
    answer_cell.fill = PatternFill("solid", fgColor="FDF1EA")

ws.auto_filter.ref = f"A2:L{2 + len(rows)}"
ws.sheet_view.zoomScale = 85

# Validation for the type column (extra blank rows keep the format friendly).
type_validation = DataValidation(
    type="list",
    formula1='"单选,多选,判断"',
    allow_blank=True,
    showDropDown=False,
)
type_validation.error = "请填写：单选 / 多选 / 判断"
type_validation.errorTitle = "题型不符合约定"
type_validation.prompt = "填写 单选、多选 或 判断"
type_validation.promptTitle = "题型"
ws.add_data_validation(type_validation)
type_validation.add("A3:A500")

difficulty_validation = DataValidation(
    type="list",
    formula1='"简单,中等,困难"',
    allow_blank=True,
    showDropDown=False,
)
ws.add_data_validation(difficulty_validation)
difficulty_validation.add("L3:L500")

score_validation = DataValidation(
    type="decimal",
    operator="greaterThan",
    formula1="0",
    allow_blank=True,
)
score_validation.error = "分值必须为正数，留空默认按 1 分计算。"
score_validation.errorTitle = "分值无效"
ws.add_data_validation(score_validation)
score_validation.add("J3:J500")

# ---------------------------------------------------------------------------
# Instructions sheet
# ---------------------------------------------------------------------------
help_ws = wb.create_sheet("填写说明")
help_ws.sheet_properties.tabColor = "DD6B43"
help_ws.sheet_view.showGridLines = False
help_ws.column_dimensions["A"].width = 16
help_ws.column_dimensions["B"].width = 16
help_ws.column_dimensions["C"].width = 64
help_ws.column_dimensions["D"].width = 34

help_ws.merge_cells("A1:D1")
help_ws["A1"] = "小点刷题 · Excel 题库填写说明"
help_ws["A1"].font = Font(name="微软雅黑", size=18, bold=True, color=CREAM)
help_ws["A1"].fill = PatternFill("solid", fgColor=BRAND)
help_ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
help_ws.row_dimensions[1].height = 46

help_ws.merge_cells("A2:D2")
help_ws["A2"] = "只读取第一张工作表「题库」。表头需保留，一行填写一道题；列的顺序不限，解析、分值、分类、难度均可留空。"
help_ws["A2"].font = Font(name="微软雅黑", size=9.5, color=MUTED)
help_ws["A2"].fill = PatternFill("solid", fgColor=CREAM_2)
help_ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
help_ws.row_dimensions[2].height = 34

header_row = 4
for col, value in enumerate(["字段", "是否必填", "填写规则", "示例"], start=1):
    cell = help_ws.cell(row=header_row, column=col, value=value)
    cell.font = Font(name="微软雅黑", size=10, bold=True, color=CREAM)
    cell.fill = PatternFill("solid", fgColor=BRAND_DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = cell_border
help_ws.row_dimensions[header_row].height = 26

instructions = [
    ("题型", "必填", "填写「单选」「多选」或「判断」。", "单选"),
    ("题干", "必填", "题目正文，支持换行和较长文本。", "HTTP 404 表示什么？"),
    ("选项A / B", "必填*", "每列一个选项。单选、多选至少填写 A、B 两个选项；判断题可留空。", "请求成功"),
    ("选项C…Z", "可选", "需要更多选项时按顺序增加列，留空的选项会被忽略。", "资源不存在"),
    ("正确答案", "必填", "单选填一个字母；多选填多个字母；判断填「正确」或「错误」。也兼容 对/错、√/×、T/F、A/B。", "ABC 或 正确"),
    ("解析", "可选", "提交或交卷后展示，支持换行。", "404 表示服务器找不到资源。"),
    ("分值", "可选", "正数，缺省按 1 分计算；考试按此累加满分。", "5"),
    ("分类", "可选", "章节或知识分类，缺省归为「未分类」。", "计算机基础"),
    ("难度", "可选", "仅作展示，可填写 简单 / 中等 / 困难。", "简单"),
]

for row_index, (field, required, rule, example) in enumerate(instructions, start=header_row + 1):
    help_ws.row_dimensions[row_index].height = 48 if field == "正确答案" else 35
    values = [field, required, rule, example]
    for col, value in enumerate(values, start=1):
        cell = help_ws.cell(row=row_index, column=col, value=value)
        cell.font = Font(
            name="微软雅黑",
            size=9.5,
            color="1F2D27",
            bold=(col == 1),
        )
        cell.alignment = Alignment(
            horizontal="center" if col == 2 else "left",
            vertical="center",
            wrap_text=True,
        )
        cell.border = cell_border
        if required == "必填":
            cell.fill = PatternFill("solid", fgColor="FDF1EA")
        elif col == 2:
            cell.fill = PatternFill("solid", fgColor=CREAM_2)

notes_start = header_row + len(instructions) + 3
help_ws.merge_cells(start_row=notes_start, start_column=1, end_row=notes_start, end_column=4)
notes_title = help_ws.cell(row=notes_start, column=1, value="导入规则")
notes_title.font = Font(name="微软雅黑", size=11, bold=True, color=BRAND)
notes_title.alignment = Alignment(horizontal="left", vertical="center")
help_ws.row_dimensions[notes_start].height = 28

notes = [
    "1. 支持 .xlsx 文件，不支持旧版 .xls；文件大小不超过 20MB。",
    "2. 表头需出现在第一张工作表的前 15 行内，程序会自动识别。",
    "3. 导入时会逐行校验；格式有误的行会在导入窗口中列出并跳过，不影响其他题目。",
    "4. 至少需要一道有效题目。导入成功后会替换当前题库，练习记录仍会保留。",
    "5. 判断题的正确答案只接受“正确 / 错误”类表达，选项列无需填写。",
    "6. 多选题答案写成 ABC 或 A,B,C 均可，顺序不影响判题。",
]
for offset, text in enumerate(notes, start=1):
    row_index = notes_start + offset
    help_ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=4)
    cell = help_ws.cell(row=row_index, column=1, value=text)
    cell.font = Font(name="微软雅黑", size=9, color=MUTED)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    cell.fill = PatternFill("solid", fgColor=SOFT_GREEN if offset % 2 else CREAM_2)
    cell.border = Border(bottom=Side(style="hair", color=LINE))
    help_ws.row_dimensions[row_index].height = 27

# Comparison section
compare_start = notes_start + len(notes) + 3
help_ws.merge_cells(start_row=compare_start, start_column=1, end_row=compare_start, end_column=4)
compare_title = help_ws.cell(row=compare_start, column=1, value="三种学习模式")
compare_title.font = Font(name="微软雅黑", size=11, bold=True, color=BRAND)
compare_title.alignment = Alignment(horizontal="left", vertical="center")

compare_headers = ["模式", "答题反馈", "计时", "结果统计"]
for col, value in enumerate(compare_headers, start=1):
    cell = help_ws.cell(row=compare_start + 1, column=col, value=value)
    cell.font = Font(name="微软雅黑", size=9.5, bold=True, color=CREAM)
    cell.fill = PatternFill("solid", fgColor=BRAND_DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = cell_border

comparisons = [
    ("做题模式", "每题提交后立即显示对错与解析", "不计时", "正确率、答对/错数量、错题解析"),
    ("背题模式", "答案与解析直接展示", "不计时", "不评分，只记录浏览数量与用时"),
    ("考试模式", "作答时不提示，交卷后统一评分", "限时或无限时", "总分、正确率、用时、错题解析"),
]
for row_offset, row_data in enumerate(comparisons, start=1):
    row_index = compare_start + 1 + row_offset
    for col, value in enumerate(row_data, start=1):
        cell = help_ws.cell(row=row_index, column=col, value=value)
        cell.font = Font(name="微软雅黑", size=9, color="1F2D27", bold=(col == 1))
        cell.alignment = Alignment(horizontal="left" if col > 1 else "center", vertical="center", wrap_text=True)
        cell.border = cell_border
        if col == 2:
            cell.fill = PatternFill("solid", fgColor=CREAM_2)
    help_ws.row_dimensions[row_index].height = 34

help_ws.sheet_view.zoomScale = 100

wb.save(OUT)
print(f"Saved {OUT} ({OUT.stat().st_size} bytes)")
